#!/usr/bin/env python3
"""
Monta a planilha de validação para a equipe de faturamento do HUB (Rebeca),
a partir de uma rodada da etapa 4 já gravada. Não roda o pipeline nem usa LLM.

Uso (a partir da raiz do projeto, com o venv ativado):
    python src/analise/montar_planilha_rebeca.py data/data-completo.xlsx --rotulo gemma27_v3
    python src/analise/montar_planilha_rebeca.py data/data-completo.xlsx --rotulo gemma27_v3 --grupo teste

Precisa de data/avaliacao/<grupo>_entrada.json e <grupo>_gabarito.json, que a
etapa 4 grava ao rodar com --grupo (padrão aqui: todos).

Só entram os códigos do escopo do estudo (os mesmos da tela de conferência).
Para cada evolução, o que o sistema sugeriu é comparado com o que o auditor
marcou, e a planilha separa três grupos:

  1. Marcados pelo auditor e não sugeridos. A equipe responde se está escrito
     na evolução (erro do sistema) ou não (inferido pelo auditor, fora do
     alcance de quem lê só a evolução). Diárias já vêm marcadas "Não".
  2. Acertos. Sistema e auditor concordam. Só informação, sem pergunta.
  3. Sugeridos pelo sistema e não marcados. A equipe responde se foi realizado
     e deveria ser cobrado (item que passou batido) ou não (erro do sistema).

Saída: data/avaliacao/planilha_rebeca_<rotulo>_<grupo>.xlsx. Tem texto real de
evolução: fica fora do git, vai para a equipe do HUB por canal institucional e
não deve ser copiada para outro lugar. O script só imprime contagens.
"""
import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[2]
PASTA_DADOS = RAIZ / "data" / "avaliacao"

# Mesmos códigos de CODIGOS_DO_ESCOPO em frontend/src/lib/itens.ts.
ESCOPO = {
    "0309010047": "NUTRIÇÃO ENTERAL EM ADULTO",
    "0401010015": "CURATIVO GRAU II C/ OU S/ DEBRIDAMENTO",
    "0211080020": "GASOMETRIA",
    "0302040021": "ATENDIMENTO FISIOTERAPÊUTICO EM PACIENTE COM TRANSTORNO RESPIRATÓRIO SEM COMPLICAÇÕES SISTÊMICAS",
    "0302050027": "ATENDIMENTO FISIOTERAPÊUTICO NAS ALTERAÇÕES MOTORAS",
    "0301100071": "CUIDADOS C/ TRAQUEOSTOMIA",
    "0305010131": "HEMODIÁLISE P/ PACIENTES RENAIS AGUDOS / CRÔNICOS AGUDIZADOS S/ TRATAMENTO DIALÍTICO INICIADO",
    "0301100055": "CATETERISMO VESICAL DE DEMORA",
    "0309010101": "PASSAGEM DE SONDA NASO ENTERICA (INCLUI SONDA)",
    "0211020036": "ELETROCARDIOGRAMA",
    "0214010015": "GLICEMIA CAPILAR",
    "0802010083": "DIÁRIA DE UNIDADE DE TERAPIA INTENSIVA ADULTO (UTI II)",
    "0301010170": "CONSULTA/AVALIAÇÃO EM PACIENTE INTERNADO",
    "0301010048": "CONSULTA DE PROFISSIONAIS DE NÍVEL SUPERIOR NA ATENÇÃO ESPECIALIZADA (EXCETO MÉDICO)",
    "0201020041": "COLETA DE MATERIAL PARA EXAME LABORATORIAL",
    "0301100144": "OXIGENOTERAPIA POR DIA",
}

# Códigos que o auditor lança pela permanência do paciente, sem registro na evolução.
SEM_EVIDENCIA_NO_TEXTO = {"0802010083": "Não (contada pela permanência na UTI)"}

# Termos principais de cada regra (os mesmos de TERMOS_REGRA no front), para
# achar o trecho onde o sistema viu o item.
TERMOS_REGRA = {
    "0401010015": r"\bcurativos?\b|\bclorexidina\b|\boclu(o|sao|ido)\b",
    "0302040021": r"\bfisioterapia respiratoria\b|\bmonitorizacao ventilatoria\b",
    "0302050027": r"\bfisioterapia motora\b|\bcinesioterapia\b|\bsedestacao\b|\bmobilizacao\b",
    "0211080020": r"\bgasometria\b|\bgas[ao]\b|\bpao2\b|\bpaco2\b|\bhco3\b",
    "0305010131": r"\bhd\b|\b(hemo)?dialise\b|\buf\b|\btrs\b",
    "0301100071": r"\btqt\b|\btraqueostomi\w*",
    "0309010047": r"\bsne\b|\bgtt\b|\bdieta enteral\b|\btne\b|\bterapia nutricional\b",
    "0301100144": r"\bvm\b|\bventilacao mecanica\b|\bfio2\b|\bcateter nasal\b|\btot\b|\btqt\b",
}

LIMITE_CELULA = 32000  # o Excel aceita até 32.767 caracteres por célula


def so_digitos(c) -> str:
    return re.sub(r"\D", "", str(c or ""))


def formatar(d: str) -> str:
    return f"{d[0:2]}.{d[2:4]}.{d[4:6]}.{d[6:9]}-{d[9]}" if len(d) == 10 else d


def normalizar(texto: str) -> str:
    """Minúsculas e sem acento, com o mesmo comprimento do original."""
    saida = []
    for ch in texto:
        base = "".join(c for c in unicodedata.normalize("NFKD", ch) if not unicodedata.combining(c)) or ch
        saida.append(base[0].lower())
    return "".join(saida)


def recorte(texto: str, inicio: int, fim: int, margem: int = 90) -> str:
    """Trecho em volta da posição, cortado em espaço, numa linha só."""
    a = max(0, inicio - margem)
    b = min(len(texto), fim + margem)
    while a > 0 and not texto[a - 1].isspace():
        a -= 1
    while b < len(texto) and not texto[b].isspace():
        b += 1
    trecho = " ".join(texto[a:b].split())
    return ("… " if a > 0 else "") + trecho + (" …" if b < len(texto) else "")


def achar_trecho(texto: str, codigo: str, origem: str, nivel: str) -> str:
    if not texto:
        return ""
    if nivel == "regra_documento":
        primeira = next((l.strip() for l in texto.splitlines() if l.strip()), "")
        return f"Tipo da nota (cabeçalho): {primeira[:200]}"
    norm = normalizar(texto)
    padroes = []
    if origem and not origem.startswith("["):
        t = re.escape(normalizar(origem.strip()))
        padroes.append(rf"(?<![a-z0-9]){t}(?![a-z0-9])")
    if codigo in TERMOS_REGRA:
        padroes.append(TERMOS_REGRA[codigo])
    for p in padroes:
        m = re.search(p, norm)
        if m:
            return recorte(texto, m.start(), m.end())
    return "(trecho não localizado automaticamente; ver o texto na aba Evoluções)"


def origem_legivel(nivel: str, origem: str) -> str:
    if nivel == "regra_documento":
        return "Tipo da nota"
    if nivel == "regra_texto":
        return "Regra de texto"
    if nivel == "regra_texto_confirmada":
        return "Regra de texto confirmada pela IA"
    return f'IA encontrou "{origem}"' if origem else "IA"


def descricoes_da_planilha(caminho: Path) -> dict:
    df = pd.read_excel(caminho, dtype=str)
    df.columns = [c.strip() for c in df.columns]
    saida = {}
    for v in df.get("itens_faturaveis", pd.Series(dtype=str)).dropna():
        for linha in str(v).splitlines():
            m = re.match(r"\s*(\d{2}\.\d{2}\.\d{2}\.\d{3}-\d)\s*-\s*(.+)", linha)
            if m:
                saida.setdefault(so_digitos(m.group(1)), m.group(2).strip())
    return saida


def ler_json(nome: str):
    arq = PASTA_DADOS / nome
    if not arq.exists():
        sys.exit(f"Não encontrei {arq}. Rode antes a etapa 4 com o --grupo correspondente.")
    return json.loads(arq.read_text(encoding="utf-8"))


# ── formatação da planilha ─────────────────────────────────────────────────

CAB = PatternFill("solid", fgColor="0B6B6F")
RESPOSTA = PatternFill("solid", fgColor="FFF4C2")
PREENCHIDO = PatternFill("solid", fgColor="E7F3EC")


def aba(livro, titulo, cabecalho, larguras):
    ws = livro.create_sheet(titulo)
    ws.append(cabecalho)
    for i, larg in enumerate(larguras, start=1):
        cel = ws.cell(row=1, column=i)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = CAB
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cel.column_letter].width = larg
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    return ws


def quebrar(ws, colunas):
    for linha in ws.iter_rows(min_row=2):
        for cel in linha:
            cel.alignment = Alignment(wrap_text=cel.column_letter in colunas, vertical="top")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha")
    ap.add_argument("--rotulo", required=True)
    ap.add_argument("--grupo", choices=["teste", "treino", "todos"], default="todos")
    args = ap.parse_args()

    entrada = {p["id"]: p for p in ler_json(f"{args.grupo}_entrada.json")}
    gabarito = ler_json(f"{args.grupo}_gabarito.json")
    arq_res = PASTA_DADOS / f"resultados_{args.rotulo}.jsonl"
    if not arq_res.exists():
        sys.exit(f"Não encontrei {arq_res}.")
    resultados = {}
    for linha in arq_res.read_text(encoding="utf-8").splitlines():
        if linha.strip():
            r = json.loads(linha)
            if "erro" not in r and r["id"] in entrada:
                resultados[r["id"]] = r

    desc_planilha = descricoes_da_planilha(Path(args.planilha))

    def descricao(d):
        return desc_planilha.get(d) or ESCOPO.get(d, "")

    ids = sorted(resultados, key=lambda i: (len(i), i))
    g1, g2, g3 = [], [], []
    for i in ids:
        r = resultados[i]
        texto = entrada[i].get("texto", "")
        auditor = gabarito.get(i, {}).get("anotador", "")
        ouro = {so_digitos(c) for c in gabarito.get(i, {}).get("codigos", [])} & ESCOPO.keys()
        sugeridos = {}
        for c in r.get("codigos", []):
            d = so_digitos(c.get("codigo"))
            if d in ESCOPO and d not in sugeridos:
                sugeridos[d] = c
        conferir = {so_digitos(c) for c in r.get("candidatos", [])}

        for d in sorted(ouro - sugeridos.keys()):
            g1.append([i, auditor, formatar(d), descricao(d),
                       "Sim" if d in conferir else "Não",
                       SEM_EVIDENCIA_NO_TEXTO.get(d, ""), "", ""])
        for d in sorted(ouro & sugeridos.keys()):
            c = sugeridos[d]
            g2.append([i, auditor, formatar(d), descricao(d),
                       origem_legivel(c.get("nivel", ""), c.get("origem", "")),
                       achar_trecho(texto, d, c.get("origem", ""), c.get("nivel", ""))])
        for d in sorted(sugeridos.keys() - ouro):
            c = sugeridos[d]
            g3.append([i, auditor, formatar(d), descricao(d),
                       origem_legivel(c.get("nivel", ""), c.get("origem", "")),
                       achar_trecho(texto, d, c.get("origem", ""), c.get("nivel", "")), "", ""])

    livro = Workbook()
    leia = livro.active
    leia.title = "Leia primeiro"
    instrucoes = [
        ["Validação do sistema de sugestão de itens de faturamento"],
        [],
        [f"{len(ids)} evoluções de UTI. Para cada uma, comparamos o que o sistema sugeriu com o que o auditor marcou."],
        ["O sistema lê uma evolução por vez e só sugere o que está escrito nela. Ele não junta evoluções do mesmo paciente."],
        ["Entram só os itens avaliados no estudo (diária, consultas, curativo, fisioterapia, gasometria, hemodiálise, oxigenoterapia etc.)."],
        [],
        ["Aba 1. Marcados e não sugeridos", f"{len(g1)} itens. O auditor marcou e o sistema não sugeriu."],
        ["", "Pergunta: está escrito nesta evolução? Se sim, cole o trecho. Se não, o auditor inferiu por outra informação (experiência, outras evoluções, permanência)."],
        ["", "As diárias já vêm preenchidas com Não, porque são contadas pela permanência. Corrija se discordar."],
        ["", "A coluna 'Ficou em Para conferir' indica quando o sistema viu indícios mas não teve certeza."],
        ["Aba 2. Acertos", f"{len(g2)} itens. Sistema e auditor concordam. Não precisa responder nada."],
        ["Aba 3. Sugeridos a mais", f"{len(g3)} itens. O sistema sugeriu e o auditor não marcou."],
        ["", "Pergunta: foi realizado e deveria ser cobrado? Sim indica item que passou na auditoria. Não indica erro do sistema."],
        ["Aba Evoluções", "Texto completo de cada evolução, para consulta pelo número."],
        [],
        ["Preencha só as colunas amarelas. Use a lista Sim ou Não. A coluna Observação é livre."],
    ]
    for linha in instrucoes:
        leia.append(linha)
    leia["A1"].font = Font(bold=True, size=14)
    leia.column_dimensions["A"].width = 34
    leia.column_dimensions["B"].width = 120
    for linha in leia.iter_rows():
        for cel in linha:
            cel.alignment = Alignment(wrap_text=True, vertical="top")

    sim_nao = '"Sim,Não"'

    ws1 = aba(livro, "1. Marcados e não sugeridos",
              ["Evolução", "Auditor", "Código", "Procedimento", "Ficou em Para conferir",
               "Está escrito nesta evolução? (Sim/Não)", "Se sim, cole o trecho", "Observação"],
              [10, 8, 15, 50, 13, 22, 50, 30])
    for linha in g1:
        ws1.append(linha)
    dv1 = DataValidation(type="list", formula1='"Sim,Não,Não (contada pela permanência na UTI)"', allow_blank=True)
    ws1.add_data_validation(dv1)
    if g1:
        dv1.add(f"F2:F{len(g1) + 1}")
    for n in range(2, len(g1) + 2):
        for col in "FGH":
            ws1[f"{col}{n}"].fill = PREENCHIDO if (col == "F" and ws1[f"F{n}"].value) else RESPOSTA
    quebrar(ws1, "DGH")

    ws2 = aba(livro, "2. Acertos",
              ["Evolução", "Auditor", "Código", "Procedimento", "Como o sistema achou", "Trecho"],
              [10, 8, 15, 50, 30, 80])
    for linha in g2:
        ws2.append(linha)
    quebrar(ws2, "DEF")

    ws3 = aba(livro, "3. Sugeridos a mais",
              ["Evolução", "Auditor", "Código", "Procedimento", "Como o sistema achou", "Trecho",
               "Foi realizado e deveria ser cobrado? (Sim/Não)", "Observação"],
              [10, 8, 15, 50, 30, 70, 22, 30])
    for linha in g3:
        ws3.append(linha)
    dv3 = DataValidation(type="list", formula1=sim_nao, allow_blank=True)
    ws3.add_data_validation(dv3)
    if g3:
        dv3.add(f"G2:G{len(g3) + 1}")
    for n in range(2, len(g3) + 2):
        for col in "GH":
            ws3[f"{col}{n}"].fill = RESPOSTA
    quebrar(ws3, "DEFH")

    ws4 = aba(livro, "Evoluções", ["Evolução", "Auditor", "Texto"], [10, 8, 140])
    cortados = 0
    for i in ids:
        texto = entrada[i].get("texto", "")
        if len(texto) > LIMITE_CELULA:
            texto = texto[:LIMITE_CELULA] + "\n[texto cortado: passou do limite de uma célula do Excel]"
            cortados += 1
        ws4.append([i, gabarito.get(i, {}).get("anotador", ""), texto])
    quebrar(ws4, "C")

    for ws in (ws1, ws2, ws3, ws4):
        ws.auto_filter.ref = ws.dimensions

    destino = PASTA_DADOS / f"planilha_rebeca_{args.rotulo}_{args.grupo}.xlsx"
    livro.save(destino)
    print(f"Evoluções: {len(ids)} (sem resultado ou com erro: {len(entrada) - len(ids)})")
    print(f"1. Marcados e não sugeridos: {len(g1)} "
          f"(já preenchidos como diária: {sum(1 for l in g1 if l[5])})")
    print(f"2. Acertos: {len(g2)}")
    print(f"3. Sugeridos a mais: {len(g3)} "
          f"(sem trecho localizado: {sum(1 for l in g3 if l[5].startswith('(trecho'))})")
    if cortados:
        print(f"Evoluções com texto cortado na célula: {cortados}")
    print(f"Planilha gravada em {destino}")


if __name__ == "__main__":
    main()
